#
# GAMS - General Algebraic Modeling System Python API
#
# Copyright (c) 2017-2026 GAMS Development Corp. <support@gams.com>
# Copyright (c) 2017-2026 GAMS Software GmbH <support@gams.com>
#
# Permission is hereby granted, free of charge, to any person obtaining a copy
# of this software and associated documentation files (the "Software"), to deal
# in the Software without restriction, including without limitation the rights
# to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
# copies of the Software, and to permit persons to whom the Software is
# furnished to do so, subject to the following conditions:
#
# The above copyright notice and this permission notice shall be included in all
# copies or substantial portions of the Software.
#
# THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
# IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
# FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
# AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
# LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
# OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
# SOFTWARE.
#

import warnings
from copy import copy
import os
import sys
import datetime
from gams.connect.agents._excel import ExcelAgent
from gams.connect.agents._excel import Workbook
from gams.connect.connectvalidator import ConnectValidator
from gams.core.gdx import GMS_SV_UNDEF
import gams.transfer as gt
from gams.transfer.syms._methods.tables import (
    _assert_axes_no_nans,
    _get_implied_dimension_from_axes,
    _flatten_and_convert,
)
import numpy as np
from openpyxl.utils.cell import column_index_from_string
from pandas.api.types import is_datetime64_any_dtype as is_datetime
import pandas as pd


class ExcelReader(ExcelAgent):
    _index_parameter_map = {
        "rdim": "rowDimension",
        "rowdimension": "rowDimension",
        "cdim": "columnDimension",
        "columndimension": "columnDimension",
        "skipempty": "skipEmpty",
        "se": "skipEmpty",
        "ignoretext": "ignoreText",
        "automerge": "autoMerge",
        "ignorerows": "ignoreRows",
        "ignorecolumns": "ignoreColumns",
        "mergedcells": "mergedCells",
    }

    def __init__(self, cdb, inst, agent_index):
        super().__init__(cdb, inst, agent_index)
        self._wb = None
        self._parse_options(self._inst)
        if os.path.splitext(self._file)[1] in [".xls"]:
            self._connect_error("The ExcelReader does not support .xls files.")

    def _parse_options(self, inst):
        inst["file"] = os.path.abspath(inst["file"])
        self._file = inst["file"]
        self._cdim = inst["columnDimension"]
        self._rdim = inst["rowDimension"]
        self._sym_type = inst["type"]
        self._symbols = self._dict_get(inst, "symbols", [])
        self._merged_cells = inst["mergedCells"]
        self._skip_empty = inst["skipEmpty"]
        self._value_subs = inst["valueSubstitutions"]
        self._index_subs = inst["indexSubstitutions"]
        self._index = inst["index"]
        self._auto_merge = inst["autoMerge"]
        self._ignore_text = inst["ignoreText"]
        self._trace = inst["trace"]
        self._is_xlsb = self._file.endswith(".xlsb")
        if self._is_xlsb:
            if not sys.platform.startswith("win"):
                self._connect_error(
                    f"Excel binary files (.xlsb) are supported on Windows only."
                )
            self._engine = "xlwings"
        else:
            self._engine = "openpyxl"

    def _apply_skip_empty(self, dim, idx, skip_empty):
        stop = None
        count = 0
        if dim > 0 and skip_empty > -1:
            for i in range(idx.shape[1]):
                if (np.array(idx[:, i] == None)).all():
                    count += 1
                else:
                    count = 0
                if count > skip_empty:
                    stop = i - skip_empty
                    break
        return stop

    def _create_index(self, dim, idx):
        if dim > 1:
            return pd.MultiIndex.from_arrays(idx)
        else:
            return idx.flatten()

    def _remove_missing_index(self, values, rdim, cdim, row_idx, col_idx, method):
        def _keep_list(idx):
            keep = list(range(idx.shape[1]))
            for i in reversed(range(idx.shape[1])):
                if method(
                    v is None or v != v for v in idx[:, i]
                ):  # drop None and float('nan') records
                    del keep[i]
            return keep

        if rdim > 0:
            keep = _keep_list(row_idx)
            row_idx = row_idx[:, keep]
            values = values[keep]

        if cdim > 0:
            keep = _keep_list(col_idx)
            col_idx = col_idx[:, keep]
            values = values[:, keep]

        return values, row_idx, col_idx

    def _apply_auto_merge(self, idx, dim):
        last_label = [None] * dim
        for i in range(idx.shape[1]):
            if any(idx[:, i] != None):
                for j in range(idx.shape[0]):
                    if idx[j, i] is None:
                        idx[j, i] = last_label[j]
                last_label = idx[:, i]
        return idx

    def _create_dataframe(self, col_idx, row_idx, values, rdim, cdim):
        # create column and row index used for DataFrame
        col_idx = self._create_index(cdim, col_idx)
        row_idx = self._create_index(rdim, row_idx)

        if cdim == rdim == 0:
            df = pd.DataFrame(values.flatten())
        elif cdim == 0:
            if values.size == 0:
                return pd.DataFrame([np.nan] * len(row_idx), index=row_idx)
            values = values[:, 0]
            df = pd.DataFrame(values.flatten(), index=row_idx)
        elif rdim == 0:
            if values.size == 0:
                return pd.DataFrame([np.nan] * len(col_idx), index=col_idx)
            values = values[0, :]
            df = pd.DataFrame(values.flatten(), index=col_idx)
        else:
            df = pd.DataFrame(values, index=row_idx, columns=col_idx)
        return df

    def _resolve_merged_cells(self, sheet, data):
        # TODO: do this only on the used range for better performance
        if self._engine == "xlwings":
            ## Windows and macOS compatible approach
            #last_cell = sheet.used_range.last_cell
            #used = sheet.range(sheet.range("A1"), last_cell)
            #ranges = set()
            #for row in range(used.rows.count):
            #    for col in range(used.columns.count):
            #        cell = used[row, col]
            #        merge_area = cell.merge_area
            #        if merge_area.address != cell.address:
            #            ranges.add(merge_area)

            # Windows only approach using COM
            ranges = {sheet.range(c.MergeArea.Address) for c in sheet.api.UsedRange.Cells if c.MergeCells}
            mr = []
            for rng in ranges:
                mr.append(
                    (
                        rng.column,
                        rng.row,
                        rng.column + rng.columns.count - 1,
                        rng.row + rng.rows.count - 1,
                    )
                )
        else:
            mr = [x.bounds for x in sheet.merged_cells.ranges]
        for b in mr:
            nwc, nwr, sec, ser = b
            value = data[nwr - 1][nwc - 1]
            data[nwr - 1 : ser, nwc - 1 : sec] = value
        return data

    def _convert_dates(self, df):
        for col in df.columns:
            if is_datetime(df[col]):
                df[col] = (
                    pd.DatetimeIndex(df[col]).to_julian_date()
                    - pd.Timestamp("1899-12-30").to_julian_date()
                )

        has_datetime = any(
            isinstance(x, datetime.datetime) for x in df.values.flatten()
        )
        if has_datetime:
            if hasattr(pd.DataFrame, "map"):
                df = df.map(
                    lambda x: (
                        pd.Timestamp(x).to_julian_date()
                        - pd.Timestamp("1899-12-30").to_julian_date()
                        if isinstance(x, datetime.datetime)
                        else x
                    )
                )
            else:
                df = df.applymap(
                    lambda x: (
                        pd.Timestamp(x).to_julian_date()
                        - pd.Timestamp("1899-12-30").to_julian_date()
                        if isinstance(x, datetime.datetime)
                        else x
                    )
                )

        return df

    def _index_substitutions(self, row_idx, col_idx, rdim, cdim, index_subs):
        if index_subs and rdim + cdim > 0:
            subs = copy(index_subs)
            for k, v in subs.items():
                if k is None:
                    break
                if k != k:  # check for float('nan')
                    if None not in subs.keys():
                        subs[None] = v
                    break
            if rdim > 0:
                ri_tmp = copy(row_idx)
                for k, v in subs.items():
                    # change value in array if either row_idx==k or (element is .nan (row_idx!=row_idx) and key is .nan (k!=k))
                    ri_tmp[
                        np.logical_or(
                            row_idx == k, np.logical_and(row_idx != row_idx, k != k)
                        )
                    ] = v
                row_idx = ri_tmp
            if cdim > 0:
                ci_tmp = copy(col_idx)
                for k, v in subs.items():
                    # change value in array if either col_idx==k or (element is .nan (col_idx!=col_idx) and key is .nan (k!=k))
                    ci_tmp[
                        np.logical_or(
                            col_idx == k, np.logical_and(col_idx != col_idx, k != k)
                        )
                    ] = v
                col_idx = ci_tmp
        return row_idx, col_idx
        # alternative approach, but much slower
        # vectorized_replace_rows = np.vectorize(lambda v: dict().get(v, v), otypes=[row_idx.dtype])
        # vectorized_replace_cols = np.vectorize(lambda v: dict().get(v, v), otypes=[col_idx.dtype])
        #
        # if rdim > 0:
        #    row_idx = vectorized_replace_rows(row_idx)
        # if cdim > 0:
        #    col_idx = vectorized_replace_cols(col_idx)

    def _value_substitutions(self, df, value_sub):
        if value_sub:
            # pandas-version-check
            if self._pandas_version_before(pd.__version__, "2.2"):  # pandas < 2.2.0
                df.replace(value_sub, inplace=True)
            else:  # pandas >= 2.2.0
                with pd.option_context("future.no_silent_downcasting", True):
                    df = df.replace(value_sub).infer_objects()
        return df

    def _write(self, df, sym_name, sym_type, rdim, cdim):
        if df is None or df.empty or df.isnull().all().all():
            df = None
        # pass DataFrame as Series for rdim=0 or cdim=0 to not confuse gams.transfer with dimensions
        elif (cdim == 0 and rdim != 0) or (rdim == 0 and cdim != 0):
            df = df[0]

        if sym_type == "par":
            sym = self._cdb.container.addParameter(
                sym_name,
                ["*"] * (rdim + cdim),
                records=df,
                uels_on_axes=True,
            )
            if df is not None:
                sym.dropUndef()  # drop float('nan')
                # TODO: remove this section as soon as gams.transfer supports dropping NaN values
                sym.records = self._value_substitutions(
                    sym.records, {GMS_SV_UNDEF: gt.SpecialValues.UNDEF}
                )

        else:  # set
            if df is not None:
                # TODO: remove this section as soon as gams.transfer supports dropping NaN values
                # Nan values become empty set element text and we can not drope those values after they are in the container.
                # This is the workaround to handle this
                _assert_axes_no_nans(df)
                dim = _get_implied_dimension_from_axes(df)
                if dim != rdim + cdim:
                    self._connect_error(
                        f"Dimensionality of table ({dim}) is inconsistent with set domain specification ({rdim+cdim})"
                    )
                df = _flatten_and_convert(df)
                df.dropna(inplace=True)
            sym = self._cdb.container.addSet(
                sym_name,
                ["*"] * (rdim + cdim),
                records=df,
                uels_on_axes=False,
            )

        # For symbols with None records, empty df is assigned
        self._transform_sym_none_to_empty(sym)

    def _ignore_rows(self, data, offset, ignore_rows):
        nr_rows = data.shape[0]
        r = list(
            filter(
                lambda x: x + offset + 1 not in ignore_rows,
                range(nr_rows),
            )
        )
        data = data[r]
        return data

    def _ignore_columns(self, data, offset, ignore_columns):
        nr_cols = data.shape[1]
        r = list(
            filter(
                lambda x: x + offset + 1 not in ignore_columns,
                range(nr_cols),
            )
        )
        data = data[:, r]
        return data

    def _apply_ignore_rows_columns(
        self, data, ignore_rows, ignore_columns, nw_row, nw_col, sym_name
    ):
        # apply ignoreRows
        if ignore_rows is not None:
            data = self._ignore_rows(data, nw_row, ignore_rows)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Raw data after ignoreRows ({sym_name}):\n{data}\n"
                )

        # apply ignoreColumns
        if ignore_columns is not None:
            data = self._ignore_columns(data, nw_col, ignore_columns)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Raw data after ignoreColumns ({sym_name}):\n{data}\n"
                )
        return data

    def _parse_ignore_rows(self, ignore_rows, nw_row, se_row):
        if ignore_rows is None:
            return []
        if isinstance(ignore_rows, int):
            l = [ignore_rows]
        elif isinstance(ignore_rows, str):
            l = self._parse_rows_range(ignore_rows)
        else:  # list
            # [9, "4:7", 11] -> [9, 4, 5, 6, 7, 11]
            l = [
                i
                for r in ignore_rows
                for i in (self._parse_rows_range(r) if isinstance(r, str) else [r])
            ]
        l = set(l)
        if se_row is None:
            return list(l)
        l = list(filter(lambda x: x >= nw_row and x <= se_row, l))
        return l

    def _parse_ignore_columns(self, ignore_columns, nw_col, se_col):
        if ignore_columns is None:
            return []
        if not isinstance(
            ignore_columns, list
        ):  # turn int and str values into a list first
            ignore_columns = [ignore_columns]
        l = []
        for c in ignore_columns:
            if isinstance(c, int):
                l.append(c)
            else:  # string
                if ":" in c:
                    l.extend(self._parse_columns_range(c))
                else:
                    l.append(column_index_from_string(c))
        l = set(l)
        if se_col is None:
            return list(l)
        l = list(filter(lambda x: x >= nw_col and x <= se_col, l))
        return l

    def _read_symbol(self, sym):
        sym_raw = sym.copy()
        self._update_sym_inst(sym, self._inst)

        sym["range"] = self._dict_get(sym, "range", sym["name"] + "!A1")
        sym["range"] = self.normalize_range(sym["range"])

        rdim = sym["rowDimension"]
        cdim = sym["columnDimension"]
        sym_range = sym["range"]
        sym_name = sym["name"]
        sym_type = sym["type"]
        merged_cells = sym["mergedCells"]
        value_subs = sym["valueSubstitutions"]
        index_subs = sym["indexSubstitutions"]
        skip_empty = sym["skipEmpty"]
        auto_merge = sym["autoMerge"]

        sheet, nw_col, nw_row, se_col, se_row, _ = self.parse_range(sym_range, self._wb)
        nw_only = se_col is None and se_row is None
        ignore_rows = self._parse_ignore_rows(sym["ignoreRows"], nw_row, se_row)
        ignore_columns = self._parse_ignore_columns(
            sym["ignoreColumns"], nw_col, se_col
        )
        required_rows = cdim + 1 + len(ignore_rows)
        required_cols = rdim + 1 + len(ignore_columns)
        if not nw_only:
            nr_cols = se_col - nw_col
            nr_rows = se_row - nw_row

        # handle ignoreText=infer
        if sym["type"] == "set" and sym["ignoreText"] == "infer":
            sym["ignoreText"] = False
            if rdim == 0:
                if (
                    nw_only or nr_rows < required_rows
                ):  # nw only or range without set element text
                    sym["ignoreText"] = True
            if cdim == 0:
                if (
                    nw_only or nr_cols < required_cols
                ):  # nw only or range without set element text
                    sym["ignoreText"] = True
        ignore_text = sym["ignoreText"]

        if self._trace > 0:
            self._log_instructions(
                sym, sym_raw, description=f"Read symbol >{sym['name']}<:"
            )

        self._symbols_exist_cdb(sym_name)

        # check that sets do not have dim=0
        if sym_type == "set" and rdim == 0 and cdim == 0:
            self._connect_error(
                f"Cannot read set >{sym_name}< with both >rowDimension: 0< and >columnDimension: 0<."
            )

        # check sufficient ranges
        if sym_type == "set" and ignore_text:
            if cdim == 0:
                required_cols -= 1
            elif rdim == 0:
                required_rows -= 1
        if not nw_only:
            if sym_type == "set" and not ignore_text:
                if cdim == 0 and nr_cols == required_cols - 1:
                    self._connect_error(
                        "Range and rowDimension specification does not contain set element text but ignoreText has been set to False. Adjust range or rowDimension or set >ignoreText: True<."
                    )
                if rdim == 0 and nr_rows == required_rows - 1:
                    self._connect_error(
                        "Range and columnDimension specification does not contain set element text but ignoreText has been set to False. Adjust range or columnDimension or set >ignoreText: True<."
                    )
            if nr_rows < required_rows:
                self._connect_error(
                    f"Invalid range >{sym_range}<. With columnDimension: >{cdim}< and {len(ignore_rows)} rows to be ignored, the range must include at least {required_rows} rows."
                )
            if nr_cols < required_cols:
                self._connect_error(
                    f"Invalid range >{sym_range}<. With rowDimension: >{rdim}< and {len(ignore_columns)} columns to be ignored, the range must include at least {required_cols} columns."
                )

        data = self._wb.get_sheet_data(sheet)

        if len(data) == 0:  # no data at all
            self._write(None, sym_name, sym_type, rdim, cdim)
            return

        if self._trace > 2:
            self._cdb.print_log(f"Raw data ({sym_name}) :\n{data}\n")

        if merged_cells:
            data = self._resolve_merged_cells(sheet, data)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Raw data after resolving merged cells ({sym_name}):\n{data}\n"
                )

        # shrink data to actual range
        data = data[nw_row:se_row, nw_col:se_col]
        if self._trace > 2:
            self._cdb.print_log(
                f"Raw data after shrinking to range ({sym_name}):\n{data}\n"
            )

        # apply ignoreRows and ignoreColumns
        data = self._apply_ignore_rows_columns(
            data, ignore_rows, ignore_columns, nw_row, nw_col, sym_name
        )
        if data.size == 0:
            self._write(None, sym_name, sym_type, rdim, cdim)
            return
        # if data.shape[0] < required_rows - len(ignore_rows):
        #    self._connect_error(
        #        f"Insufficient number of data rows ({sym_name}). Require at least {required_rows}, but got {data.shape[0]}."
        #    )
        # if data.shape[1] < required_cols - len(ignore_columns):
        #    self._connect_error(
        #        f"Insufficient number of data columns ({sym_name}). Require at least {required_cols}, but got {data.shape[1]}."
        #    )

        col_idx = data[:cdim, rdim:]
        row_idx = data[cdim:, :rdim].transpose()

        if self._trace > 2:
            self._cdb.print_log(f"Initial column index ({sym_name}):\n{col_idx}\n")
            self._cdb.print_log(f"Initial row index ({sym_name}):\n{row_idx}\n")

        # apply skipEmpty only for nw_only, but not for explicit ranges
        if nw_only:
            stop_col = self._apply_skip_empty(cdim, col_idx, skip_empty)
            col_idx = col_idx[:, :stop_col]
            stop_row = self._apply_skip_empty(rdim, row_idx, skip_empty)
            row_idx = row_idx[:, :stop_row]
            if self._trace > 2:
                self._cdb.print_log(
                    f"Column index after skipEmpty ({sym_name}):\n{col_idx}\n"
                )
                self._cdb.print_log(
                    f"Row index after skipEmpty ({sym_name}):\n{row_idx}\n"
                )
        else:
            stop_col = None
            stop_row = None

        if stop_col is not None:
            stop_col += rdim
        if stop_row is not None:
            stop_row += cdim

        if cdim == 0 and rdim == 0:  # handle scalars
            stop_row = 1
            stop_col = 1
        if rdim == 0:
            row_idx = np.empty((0, 0))  # dummy array for header
        if cdim == 0:
            col_idx = np.empty((0, 0))  # dummy array for header

        values = data[cdim:stop_row, rdim:stop_col]

        if self._trace > 2:
            self._cdb.print_log(f"Values {(sym_name)}: {values}\n")

        if self._engine == "xlwings":
            col_idx = self._to_int_if_whole(col_idx)
            row_idx = self._to_int_if_whole(row_idx)
            values = self._to_int_if_whole(values)

        if auto_merge:
            if cdim > 1:
                col_idx = self._apply_auto_merge(col_idx, cdim)
            if rdim > 1:
                row_idx = self._apply_auto_merge(row_idx, rdim)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Row index after autoMerge ({sym_name}):\n{row_idx}\n"
                )
                self._cdb.print_log(
                    f"Column index after autoMerge ({sym_name}):\n{col_idx}\n"
                )

        # replace all set text with empty string for ignoreText=True
        if sym_type == "set" and ignore_text:
            if values.size == 0:
                if cdim == 0:
                    values = np.empty((values.shape[0], 1), dtype=str)
                elif rdim == 0:
                    values = np.empty((1, values.shape[1]), dtype=str)
            else:
                values = np.empty_like(values, dtype=str)

        if index_subs:
            # remove all-None entries in column and row header and corresponding values
            values, row_idx, col_idx = self._remove_missing_index(
                values, rdim, cdim, row_idx, col_idx, all
            )
            row_idx, col_idx = self._index_substitutions(
                row_idx, col_idx, rdim, cdim, index_subs
            )
            if self._trace > 2:
                self._cdb.print_log(
                    f"Row index after indexSubstitutions ({sym_name}):\n{row_idx}\n"
                )
                self._cdb.print_log(
                    f"Column index after indexSubstitutions ({sym_name}):\n{col_idx}\n"
                )

        # remove any-None entries in column and row header and corresponding values
        values, row_idx, col_idx = self._remove_missing_index(
            values, rdim, cdim, row_idx, col_idx, any
        )

        if self._trace > 2:
            self._cdb.print_log(
                f"Column index before DataFrame creation ({sym_name}):\n{col_idx}\n"
            )
            self._cdb.print_log(
                f"Row index before DataFrame creation ({sym_name}):\n{row_idx}\n"
            )
            self._cdb.print_log(
                f"Values before DataFrame creation ({sym_name}):\n{values}\n"
            )

        df = self._create_dataframe(col_idx, row_idx, values, rdim, cdim)

        if self._trace > 2:
            self._cdb.print_log(f"Initial DataFrame ({sym_name}):\n{df}\n")

        df = self._convert_dates(df)

        df = self._value_substitutions(df, value_subs)
        if self._trace > 2:
            self._cdb.print_log(
                f"DataFrame after valueSubstitutions ({sym_name}):\n{df}\n"
            )

        # TODO: This is a workaround to get UNDEF to survive sym.dropNA/sym.dropUndef - remove as soon as gams.transfer supports dropping NaN values
        if sym_type == "par":
            import re

            pattern = re.compile(r"undef", re.IGNORECASE)
            # pandas-version-check
            if self._pandas_version_before(pd.__version__, "2.2"):  # pandas < 2.2.0
                df.replace(regex=pattern, value=GMS_SV_UNDEF, inplace=True)
            else:  # pandas >= 2.2.0
                with pd.option_context("future.no_silent_downcasting", True):
                    df = df.replace(regex=pattern, value=GMS_SV_UNDEF).infer_objects()
        self._write(df, sym_name, sym_type, rdim, cdim)

    def _open(self):
        read_only = not (
            any(sym["mergedCells"] for sym in self._symbols) or self._merged_cells
        )
        try:
            self._wb = Workbook(
                self._file, engine=self._engine, read_only=read_only, data_only=True
            )  # data_only=True is required to read values instead of formulas
        except PermissionError as e:
            self._connect_error(
                str(e)
                + "\nThe file may already be open and might need to be closed first."
            )

    def _read_symbols(self, symbols, validate=False):
        if validate:
            sym_schema = self._cdb.load_schema(self)["symbols"]["schema"]["schema"]
            v = ConnectValidator(sym_schema)
        for i, sym in enumerate(symbols):
            if validate:
                sym = v.validated(sym)
                if sym is None:
                    self._connect_error(
                        f"Validation of item {i} in index failed: {v.errors}"
                    )
                sym = v.normalize_of_rules(sym)
            self._read_symbol(sym)

    def _create_symbol_instructions(self, rec):
        is_symbol = not None in (rec[0], rec[1], rec[2])
        inst = {}
        if is_symbol:
            inst["type"] = rec[0].lower().strip()
            inst["name"] = rec[1].strip()
            inst["range"] = rec[2].strip()
            if self._trace > 1:
                self._cdb.print_log(
                    f"\nParsing symbol >{inst['name']}< with >type: {inst['type']}< and >range: {inst['range']}<."
                )
            if inst["type"] == "dset":
                inst["type"] = "set"
                self._cdb.print_log(
                    f"Warning: Processing unsupported >type: dset< as >type: set< for symbol >{inst['name']}<."
                )
        return inst

    def _read_from_index(self):
        symbols = self.parse_index(self._index, self._wb, self._index_parameter_map)

        # reopen the file with read_only=False if required
        if not self._merged_cells:
            read_only = not any(
                self._dict_get(sym, "mergedCells", self._merged_cells)
                for sym in symbols
            )
            if not read_only and not self._is_xlsb:
                self._wb.close()
                try:
                    self._wb = Workbook(
                        self._file,
                        engine=self._engine,
                        read_only=read_only,
                        data_only=True,
                    )  # data_only=True is required to read values instead of formulas
                except PermissionError as e:
                    self._connect_error(
                        str(e)
                        + "\nThe file may already be open and might need to be closed first."
                    )
        self._read_symbols(symbols, True)

    def execute(self):
        if self._trace > 0:
            self._log_instructions(self._inst, self._inst_raw)
            self._describe_container(self._cdb.container, "Connect Container (before):")

        try:
            self._open()

            if self._index:
                self._read_from_index()
            else:
                self._read_symbols(self._symbols)
            if self._trace > 2:
                for name, sym in self._cdb.container.data.items():
                    self._cdb.print_log(
                        f"Connect Container symbol >{name}<:\n {sym.records}\n"
                    )
            if self._trace > 0:
                self._describe_container(
                    self._cdb.container, "Connect Container (after):"
                )
        finally:
            if self._wb is not None:
                self._wb.close()
