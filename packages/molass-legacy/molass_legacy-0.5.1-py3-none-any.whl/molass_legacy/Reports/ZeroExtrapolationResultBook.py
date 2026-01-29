"""

    ZeroExtrapolationResultBook.py

    Copyright (c) 2016-2025, SAXS Team, KEK-PF

"""
import os
import numpy                as np
import logging
# from openpyxl               import Workbook
from openpyxl.chart         import LineChart, ScatterChart, Reference, Series
from openpyxl.chart.series_factory  import SeriesFactory
from openpyxl.chart.error_bar       import ErrorBars
from openpyxl.chart.data_source     import NumDataSource, NumData, NumVal
from openpyxl.chart.layout  import Layout, ManualLayout
from molass_legacy.KekLib.OpenPyXlUtil import save_allowing_user_reply, LINE_WIDTH
from molass_legacy.KekLib.ExceptionTracebacker   import ExceptionTracebacker
from molass_legacy.KekLib.HtmlColorNames import *
from molass_legacy.Test.TesterLogger      import write_to_tester_log
from molass_legacy._MOLASS.SerialSettings         import get_setting
from molass_legacy.Reports.DefaultFont import set_default_font
set_default_font()

CHART_WIDTH_ZX_P    = 21
CHART_HEIGHT_ZX_P   = 14
CHART_WIDTH_ZX_I    = 62.5
CHART_HEIGHT_ZX_I   = 24
PARAM_A_NAME    = 'A(q) - Scattering Intensity without Interparticle Effects (1.0 mg/ml)'
PARAM_B_NAME    = 'B(q) - Interparticle Effects'
DUMP_INPUT      = False
ADD_ERROR_BARS  = True
ATSAS_COLOR     = darkgray

"""
    borrowed from https://github.com/uskysd/openpyxl-errorbar
"""
def list2errorbars(plus, minus, errDir='y', errValType='cust'):
    "Returns ErrorBar from lists of error values"

    #Convert to list of NumVal
    numvals_plus = [NumVal(i, None, v=x) for i,x in enumerate(plus)]
    numvals_minus = [NumVal(i, None, v=x) for i,x in enumerate(minus)]

    # Convert to NumData
    nd_plus = NumData(pt=numvals_plus)
    nd_minus = NumData(pt=numvals_minus)

    # Convert to NumDataSource
    nds_plus = NumDataSource(numLit=nd_plus)
    nds_minus = NumDataSource(numLit=nd_minus)

    return ErrorBars(plus=nds_plus, minus=nds_minus, errDir=errDir, errValType=errValType)

def set_layout_small( chart, adjust=0 ):
    chart.layout = Layout(
        ManualLayout(
        x=0.1+adjust, y=0.1,
        h=0.8,  w=0.76+adjust,
        xMode="edge",
        yMode="edge",
        )
    )

def set_layout_large( chart, adjust=0, r_adjust=0 ):
    chart.layout = Layout(
        ManualLayout(
        x=0.04-adjust, y=0.1,
        h=0.8,  w=0.92+adjust+r_adjust,
        xMode="edge",
        yMode="edge",
        )
    )

def create_extrapolation_params_chart_( ws, ud, num_rows, row_offset, col_offset, j, param_name, exstra_col, no_orig_param_array, r_adjust=0, overlay=False ):
    c_ = ScatterChart()
    c_.title = param_name
    c_.style = 13
    c_.y_axis.title = 'log( I )' if j == 0 else 'Intensity'
    c_.x_axis.title = 'Q(Å⁻¹)'

    min_row = row_offset + 1
    max_row = row_offset + num_rows + 1

    xvalues = Reference(ws, min_col=1,
                            min_row=min_row + 1,    # must not include column title
                            max_row=max_row)

    if no_orig_param_array:
        min_col = col_offset + j * 2
        max_col = min_col + exstra_col
    else:
        min_col = col_offset + j * 3
        max_col = min_col + 1 + exstra_col

    if overlay:
        colors = [ red, blue ]
    else:
        primary_color = red if ud == 0 else blue
        if no_orig_param_array:
            colors = [ primary_color ]
        else:
            # TODO: test this case
            colors = [ lightgray, primary_color ]
        if exstra_col == 1:
            # colors.append( darkseagreen )
            colors.append( ATSAS_COLOR )

    for col in range( min_col, max_col+1 ):
        values = Reference(ws, min_col=col, min_row=min_row, max_row=max_row)
        series = Series(values, xvalues, title_from_data=True)
        c_.series.append(series)

    if j == 0:
        c_.y_axis.scaling.logBase = 10

    c_.width    = CHART_WIDTH_ZX_I
    c_.height   = CHART_HEIGHT_ZX_I

    # Style the lines

    for i, color in enumerate(colors):
        s_ = c_.series[i]
        # s_.marker.symbol = "circle"
        s_.marker.graphicalProperties.solidFill         = color
        s_.marker.graphicalProperties.line.solidFill    = color
        s_.graphicalProperties.line.solidFill           = color
        width = LINE_WIDTH    # in EMUs
        s_.graphicalProperties.line.width               = width

    set_layout_large( c_, r_adjust=r_adjust )

    return c_

ERROR_SCALE_FOR_AXIS = 2

def get_max( a, b, error, atsas=False, max_limit_ratio=1.5 ):
    if atsas:
        error = 0
    else:
        error *= ERROR_SCALE_FOR_AXIS
    if a is None:
        if b is None:
            return None
        else:
            return b + error
    else:
        if b is None:
            return a
        else:
            b += error
            if atsas:
                if a < b and b < a*max_limit_ratio:
                    # do not adopt too large b
                    return b
                else:
                    return a
            else:
                if a < b:
                    return b
                else:
                    return a

def get_min( a, b, error, atsas=False ):
    if atsas:
        error = 0
    else:
        error *= ERROR_SCALE_FOR_AXIS
    if a is None:
        if b is None:
            return None
        else:
            return b - error
    else:
        if b is None:
            return a
        else:
            b -= error
            if a < b:
                return a
            else:
                return b

def divide( a, b, ret ):
    if a is None:
        return ret
    else:
        return a/b

class ZeroExtrapolationResultBook:
    def __init__( self, wb, ws,
                    mjn, indeces_, c_rg_iz, qvector, data_matrix, j0,
                    ze_array, param_array, orig_param_array=None,
                    guinier_boundary=None,
                    almerge_result=None, atsas_result=None,
                    boundary_j=None, parent=None, lrf_info=None ):
        m, j, num_peaks_to_exec, peak_num_ranges = mjn
        # print( 'ZeroExtrapolationResultBook: mjn=', mjn )
        self.logger = logging.getLogger( __name__  )
        self.parent = parent
        self.doing_sec = parent.doing_sec
        self.need_bq = lrf_info.need_bq()
        self.lrf_boundary_j = lrf_info.boundary_j

        if almerge_result is None:
            exz_array       = None
            overlap_from_max  = None
        else:
            exz_array       = almerge_result.exz_array
            overlap_from_max  = almerge_result.overlap_from_max

        indeces = j0 + np.array(indeces_)

        if DUMP_INPUT and m == 0 and j == 0:
            from molass_legacy.Reports.ReportUtils import dump_zx_report_args
            dump_zx_report_args( m, j, indeces, c_rg_iz, qvector, data_matrix, ze_array, param_array, orig_param_array, exz_array )

        self.m  = m
        self.ud = j
        self.num_peaks_to_exec = num_peaks_to_exec
        self.guinier_boundary = guinier_boundary
        self.indeces = indeces
        self.wb = wb
        self.ws = ws
        num_intensity_points = param_array.shape[0]
        self.num_rows = num_intensity_points + 1
        paren = '' if num_peaks_to_exec == 1 else '(%d)' % ( m + 1 )
        range_type = get_setting('range_type')
        if self.doing_sec:
            if peak_num_ranges == 1 and range_type < 5:
                comp = 'Both-side'
            else:
                comp = 'Asc-side' if self.ud == 0 else 'Desc-side'
        else:
            comp = 'Nat-comp' if self.m == 0 else 'Unf-comp'

        ws.title = '%s Extrapolation%s' % ( comp, paren )

        ws.append( [ 'Elution №', 'Input', '_MOLASS',
                        'Input', '_MOLASS', 'ATSAS',
                        'Input', '_MOLASS', 'ATSAS',
                        'Rg-error', 'x-Rg-error', 'atsas-x-Rg-error',
                        'I(0)/Conc-error', 'x-I(0)/Conc-error', 'atsas-x-I(0)/Conc-error' ] )
        self.nx_rows = nx_rows = len( c_rg_iz[0][0] )
        self.zx_rows = zx_rows = len( c_rg_iz[0][1] )
        # print( 'nx_rows=', nx_rows, ', zx_rows=', zx_rows )
        assert( nx_rows <= zx_rows )

        self.orig_param_array = orig_param_array
        self.exz_array      = exz_array
        self.overlap_from_max = overlap_from_max
        self.boundary_j     = boundary_j

        c_vector    = c_rg_iz[0][0]
        self.c_max  = np.max( c_vector )
        zx_only_rows = zx_rows - nx_rows

        # TODO: c_rg_iz refactoring

        rg_min = None
        rg_max = None
        iz_min = None
        iz_max = None

        rg_errors   = []
        rgx_errors  = []
        atsas_rgx_errors  = []
        iz_errors   = []
        izx_errors  = []
        atsas_izx_errors = []
        if self.ud == 0:
            for i in range(zx_rows):
                if atsas_result is None or i > 0:
                    atsas_Rg        = None
                    atsas_Iz        = None
                else:
                    atsas_Rg        = atsas_result.Rg
                    # atsas_Rg_error  = atsas_result.Rg_stdev
                    # atsas_Iz should have been scaled by the max concentration
                    atsas_Iz        = divide( atsas_result.I0, self.c_max, None )
                    # atsas_Iz_error  = divide( atsas_result.I0_stdev, self.c_max, 0 )

                atsas_Rg_error  = 0   # or np.nan?
                atsas_Iz_error  = 0   # or np.nan?

                concx   = c_rg_iz[0][1][i]
                rgx_, rgx_error   = c_rg_iz[1][1][i]


                rgx_errors.append( rgx_error )
                # atsas_rgx_errors.append( atsas_Rg_error )
                atsas_rgx_errors.append( 0 )        # suppress showing errorbars for atsas
                izx_, izx_error = c_rg_iz[2][1][i]

                rg_min = get_min( rg_min, rgx_, rgx_error )
                rg_max = get_max( rg_max, rgx_, rgx_error )
                rg_min = get_min( rg_min, atsas_Rg, atsas_Rg_error, atsas=True )
                rg_max = get_max( rg_max, atsas_Rg, atsas_Rg_error, atsas=True )
                iz_min = get_min( iz_min, izx_, izx_error )
                iz_max = get_max( iz_max, izx_, izx_error )
                iz_min = get_min( iz_min, atsas_Iz, atsas_Iz_error, atsas=True )
                iz_max = get_max( iz_max, atsas_Iz, atsas_Iz_error, atsas=True )
                # self.logger.info( 'debug: ' + str( (i, rg_min, rg_max, iz_min, iz_max) ) )

                # we have changed so that iz has been already scaled by /= concx
                # izx_errors.append( izx_error/concx if izx_error is not None and concx > 0 else 0 )
                # iz_concx = izx_ / concx if izx_ is not None and concx > 0 else None
                izx_errors.append( izx_error )
                iz_concx = izx_
                atsas_izx_errors.append( atsas_Iz_error )

                if i < zx_only_rows:
                    rg_errors.append( 0 )   # or np.nan?
                    iz_errors.append( 0 )   # or np.nan?
                    conc = 0 if i == 0 else None
                    ws.append( [ None, conc, concx,
                                    None, rgx_, atsas_Rg,
                                    None, iz_concx, atsas_Iz,
                                    None, rgx_error, atsas_Rg_error,
                                    None, izx_error, atsas_Iz_error ] )
                else:
                    nxi = i - zx_only_rows
                    rg_, rg_error, atsas_Rg, atsas_Rg_error = c_rg_iz[1][0][nxi]
                    rg_errors.append( rg_error )
                    iz_, iz_error, atsas_Iz, atsas_Iz_error = c_rg_iz[2][0][nxi]
                    # iz_ has not been scaled by /= concx
                    iz_error_c = iz_error/concx if iz_error is not None and concx > 0 else 0
                    iz_errors.append( iz_error_c )
                    iz_conc = iz_ / concx if iz_ is not None and concx > 0 else None
                    if atsas_Iz is not None:
                        atsas_Iz /= concx
                        atsas_Iz_error /= concx
                    ws.append( [    indeces[nxi],
                                    c_rg_iz[0][0][nxi], concx,
                                    rg_, rgx_, atsas_Rg,
                                    iz_conc, iz_concx, atsas_Iz,
                                    rg_error, rgx_error, atsas_Rg_error,
                                    iz_error_c, izx_error, atsas_Iz_error ] )
                    rg_min = get_min( rg_min, rg_, rg_error )
                    rg_max = get_max( rg_max, rg_, rg_error )
                    iz_min = get_min( iz_min, iz_conc, iz_error_c )
                    iz_max = get_max( iz_max, iz_conc, iz_error_c )

                # self.logger.info( 'debug: ' + str( (i, rg_min, rg_max, iz_min, iz_max) ) )
        else:
            for i in range(zx_rows):
                if atsas_result is None or i < zx_rows-1:
                    atsas_Rg        = None
                    atsas_Iz        = None
                else:
                    atsas_Rg        = atsas_result.Rg
                    # atsas_Rg_error  = atsas_result.Rg_stdev
                    # atsas_Iz should have been scaled by the max concentration
                    atsas_Iz        = divide( atsas_result.I0, self.c_max, None )
                    # atsas_Iz_error  = divide( atsas_result.I0_stdev, self.c_max, 0 )

                atsas_Rg_error  = 0   # or np.nan?
                atsas_Iz_error  = 0   # or np.nan?

                concx   = c_rg_iz[0][1][i]
                rgx_, rgx_error = c_rg_iz[1][1][i]
                rgx_errors.append( rgx_error )
                # atsas_rgx_errors.append( atsas_Rg_error )
                atsas_rgx_errors.append( 0 )        # suppress showing errorbars for atsas
                izx_, izx_error = c_rg_iz[2][1][i]

                rg_min = get_min( rg_min, rgx_, rgx_error )
                rg_max = get_max( rg_max, rgx_, rgx_error )
                rg_min = get_min( rg_min, atsas_Rg, atsas_Rg_error, atsas=True )
                rg_max = get_max( rg_max, atsas_Rg, atsas_Rg_error, atsas=True )
                iz_min = get_min( iz_min, izx_, izx_error )
                iz_max = get_max( iz_max, izx_, izx_error )
                iz_min = get_min( iz_min, atsas_Iz, atsas_Iz_error, atsas=True )
                iz_max = get_max( iz_max, atsas_Iz, atsas_Iz_error, atsas=True )
                # self.logger.info( 'debug: ' + str( (i, rg_min, rg_max, iz_min, iz_max) ) )

                # we have changed so that izx has been already scaled by /= concx
                # izx_errors.append( izx_error/concx if izx_error is not None and  concx > 0 else 0 )
                # iz_concx = izx_ / concx if izx_ is not None and concx > 0 else None
                izx_errors.append( izx_error )
                iz_concx = izx_
                atsas_izx_errors.append( atsas_Iz_error )

                if i < nx_rows:
                    rg_, rg_error, atsas_Rg, atsas_Rg_error = c_rg_iz[1][0][i]
                    rg_errors.append( rg_error )
                    iz_, iz_error, atsas_Iz, atsas_Iz_error = c_rg_iz[2][0][i]
                    # iz_ has not been scaled by /= concx
                    iz_error_c = iz_error/concx if iz_error is not None and  concx > 0 else 0
                    iz_errors.append( iz_error_c )
                    iz_conc = iz_ / concx if iz_ is not None and concx > 0 else None
                    if atsas_Iz is not None:
                        atsas_Iz /= concx
                        atsas_Iz_error /= concx
                    ws.append( [    indeces[i],
                                    c_rg_iz[0][0][i], concx,
                                    rg_, rgx_, atsas_Rg,
                                    iz_conc, iz_concx, atsas_Iz,
                                    rg_error, rgx_error, atsas_Rg_error,
                                    iz_error_c, izx_error, atsas_Iz_error ] )
                    rg_min = get_min( rg_min, rg_, rg_error )
                    rg_max = get_max( rg_max, rg_, rg_error )
                    iz_min = get_min( iz_min, iz_conc, iz_error_c )
                    iz_max = get_max( iz_max, iz_conc, iz_error_c )
                else:
                    rg_errors.append( 0 )   # or np.nan?
                    iz_errors.append( 0 )   # or np.nan?
                    ws.append( [ None, None, concx,
                                    None, rgx_, atsas_Rg,
                                    None, iz_concx, atsas_Iz,
                                    None, rgx_error, atsas_Rg_error, 
                                    None, izx_error, atsas_Iz_error ] )

                # self.logger.info( 'debug: ' + str( (i, rg_min, rg_max, iz_min, iz_max) ) )

        q_colname   = [ 'Q' ]

        colnames_ = q_colname

        if orig_param_array is None:
            colnames_ += [ 'A(q)', 'ATSAS(scaled)', 'B(q)' ]
        else:
            colnames_ += [ 'original A(q)', 'A(q)', 'ATSAS(scaled)', 'original B(q)', 'B(q)' ]
        ws.append( colnames_ )

        if almerge_result is not None:
            comp_atsas_Aq_list = []

        for i in range( num_intensity_points ):
            if exz_array is None:
                atsas_val = None
            else:
                atsas_val = exz_array[i,1] / self.c_max
                # A(q) intensity values should be positive for logarithmic scaling.
                # Negative values result in COM-drived Excel errors.
                if atsas_val <= 0:
                    atsas_val = None

            # A(q) intensity values should be positive for logarithmic scaling.
            # Negative values result in COM-drived Excel errors.
            a_ = param_array[i,0]
            if a_ <= 0:
                a_ = None

            if orig_param_array is None:
                param_array_slice =  np.array( [ a_, atsas_val, param_array[i,1] ] )
            else:
                # A(q) intensity values should be positive for logarithmic scaling.
                # Negative values result in COM-drived Excel errors.
                oa_ = orig_param_array[i,0]
                if oa_ <= 0:
                    oa_ = None
                param_array_slice = np.array( [ oa_, a_, atsas_val,
                                                orig_param_array[i,1], param_array[i,1],
                                              ] )
            ws.append( list( np.hstack( [ qvector[i], param_array_slice ] ) ) )
            if almerge_result is not None:
                a_ = np.nan if a_ is None else a_
                atsas_val = np.nan if atsas_val is None else atsas_val
                comp_atsas_Aq_list.append( [ a_, atsas_val ] )

        if almerge_result is not None:
            # TODO: disable this in product release
            self.compute_diff_atsas( comp_atsas_Aq_list, qvector )

        rg_errors = [ np.array( rg_errors ), np.array( rgx_errors ), np.array( atsas_rgx_errors ) ]
        iz_errors = [ np.array( iz_errors ), np.array( izx_errors ), np.array( atsas_izx_errors ) ]

        chart_start_col = 'Q'
        self.create_concetration_chart( ws, chart_start_col+'2' )
        self.create_Rg_chart( ws, 'AB2', rg_errors, [rg_min, rg_max] )
        self.create_Iz_conc_chart( ws, 'AM2', iz_errors, [iz_min, iz_max] )
        self.create_extrapolation_params_chart( ws, zx_rows+1, 2, 0, PARAM_A_NAME, chart_start_col+'33', r_adjust=-0.02 )
        if self.need_bq:
            self.create_extrapolation_params_chart( ws, zx_rows+1, 2, 1, PARAM_B_NAME, chart_start_col+'85' )

    def compute_diff_atsas( self, comp_atsas_Aq_list, qvector ):
        if self.boundary_j is None:
            return

        j = self.boundary_j
        comp_atsas_array = np.array( comp_atsas_Aq_list[0:j] )
        k_val = comp_atsas_array[:,0]
        a_val = comp_atsas_array[:,1]
        isfinite = np.logical_and( np.isfinite(k_val), np.isfinite(a_val) )
        try:
            diff_atsas = np.sum( np.abs( (k_val - a_val)[isfinite] ) )
            max_val = np.max( a_val[isfinite] )
        except:
            diff_atsas = None
            max_val = None
        write_to_tester_log( 'ex-curves-diff=' + str( [ diff_atsas, max_val, j, qvector[j] ] ) + '\n' )

    def create_params_chart( self, ws, pos, title, y_title, x_title, y_cols, x_cols, errors=None, minmax=None ):
        f, t = y_cols
        if t-f == 1:
            c_ = LineChart()
        else:
            c_ = ScatterChart()
        c_.title = title
        c_.style = 13
        c_.y_axis.title = y_title
        c_.x_axis.title = x_title

        xvalues0 = Reference(ws, min_col=x_cols[0],
                                min_row=1 + 1,      # must not include column title
                                max_row=self.zx_rows+1)

        if t-f == 1:
            data = Reference(ws, min_col=f, max_col=t, min_row=1, max_row=self.zx_rows+1 )
            c_.add_data(data, titles_from_data=True)
            c_.set_categories(xvalues0)
            c_.width    = CHART_WIDTH_ZX_P
            c_.height   = CHART_HEIGHT_ZX_P
        else:
            assert t-f == 2
            xvalues1 = Reference(ws, min_col=x_cols[1],
                                    min_row=1 + 1,      # must not include column title
                                    max_row=self.zx_rows+1)
            for col in [f, f+1, t]:
                data = Reference(ws, min_col=col, min_row=1, max_row=self.zx_rows+1 )
                xvalues_ = xvalues1 if col == f+1 else xvalues0
                series  = Series(data, xvalues_, title_from_data=True)
                c_.series.append(series)
                c_.width    = CHART_WIDTH_ZX_P
                c_.height   = CHART_HEIGHT_ZX_P

        if minmax is not None:
            # self.logger.info( 'minmax=' + str(minmax) )
            min_, max_ = minmax
            # TODO: improve scale
            if max_ is None:
                # max_ is None for 2160227. why?
                self.logger.warning( 'Unexpected state: minmax values seem to be None in create_params_chart' )
            else:
                if max_ > 10:
                    scale = 1
                else:
                    scale = 1e-3
                c_.y_axis.scaling.min   = np.floor( min_/scale )*scale
                c_.y_axis.scaling.max   = np.ceil( max_/scale )*scale

        colors = [ green, orange ]
        if t - f == 2:
            # i.e., for Rg and Iz_con
            # TODO: better expression
            colors.append( ATSAS_COLOR )

        # print( 'errors=', errors )

        # Style the lines

        for i, color in enumerate(colors):
            s_ = c_.series[i]

            if ADD_ERROR_BARS and errors:
                plus    = errors[i]
                minus   = plus
                errorbars = list2errorbars(plus, minus, errDir='y')
                s_.errBars = errorbars

            s_.marker.symbol = "circle"
            s_.marker.graphicalProperties.solidFill         = color
            s_.marker.graphicalProperties.line.solidFill    = color
            s_.graphicalProperties.line.solidFill           = color
            s_.graphicalProperties.line.width               = LINE_WIDTH    # in EMUs

        set_layout_small( c_ )

        ws.add_chart(c_, pos)

    def create_concetration_chart( self, ws, pos ):
        self.create_params_chart( ws, pos, 'Concentration', 'Concentration', 'Elution №', [ 2, 3 ], [ 1, 1 ] )

    def create_Rg_chart( self, ws, pos, errors, minmax ):
        self.create_params_chart( ws, pos, 'Rg Extrapolated', 'Rg', 'Concentration', [ 4, 6 ], [ 2, 3 ], errors=errors, minmax=minmax )

    def create_Iz_conc_chart( self, ws, pos, errors, minmax ):
        self.create_params_chart( ws, pos, 'I(0)/Conc. Extrapolated', 'I(0)/Conc.', 'Concentration', [ 7, 9 ], [ 2, 3 ], errors=errors, minmax=minmax )

    def create_extrapolated_intensity_chart( self, ws, row_offset, col_offset, pos):
        c_ = ScatterChart()
        c_.title = "Extrapolation to Zero Concentration"
        c_.style = 13
        c_.y_axis.title = 'Intensity / Concentration' if col_offset == 0 else 'Intensity'
        c_.x_axis.title = 'Q(Å⁻¹)'

        min_row = row_offset + 1
        max_row = row_offset + self.num_rows + 1

        xvalues = Reference(ws, min_col=1,
                                min_row=min_row + 1,    # must not include column title
                                max_row=max_row)

        for col in range(col_offset+2, col_offset+self.nx_rows+self.zx_rows+2):
            values  = Reference(ws, min_col=col, min_row=min_row, max_row=max_row)
            series  = Series(values, xvalues, title_from_data=True)
            c_.series.append(series)

        c_.width    = CHART_WIDTH_ZX_I
        c_.height   = CHART_HEIGHT_ZX_I

        orange_list = [ orange ] * ( self.zx_rows - 1 )
        colors_x = ( [ red ] + orange_list ) if self.ud == 0 else ( orange_list + [ red ] )
        colors = [ green ] * self.nx_rows + colors_x
        # Style the lines

        for i, color in enumerate(colors):
            s_ = c_.series[i]
            # s_.marker.symbol = "circle"
            s_.marker.graphicalProperties.solidFill         = color
            s_.marker.graphicalProperties.line.solidFill    = color
            s_.graphicalProperties.line.solidFill           = color
            width = LINE_WIDTH    # in EMUs
            if color == red:
                width *= 2
            s_.graphicalProperties.line.width               = width

        set_layout_large( c_ )

        ws.add_chart(c_, pos)

    def create_extrapolation_params_chart( self, ws, row_offset, col_offset, j, param_name, pos, r_adjust=0):

        atsas_col = 0 if ( self.exz_array is None or j == 1 ) else 1
        no_orig_param_array = self.orig_param_array is None

        chart = create_extrapolation_params_chart_( ws, self.ud, self.num_rows, row_offset, col_offset, j, param_name, atsas_col, no_orig_param_array, r_adjust )

        ws.add_chart(chart, pos)

    def add_format_setting( self, book_file ):
        from .ZeroExExcelFormatter import ZeroExReportArgs
        axis_direction_desc = get_setting( 'axis_direction_desc' )
        args = ZeroExReportArgs(self.ws.title, self, book_file, axis_direction_desc)    # make it picklable

        if self.parent.more_multicore:
            self.add_format_setting_more_multicore(args)
        else:
            self.add_format_setting_less_multicore(args)

    def add_format_setting_more_multicore(self, args):
        self.parent.teller.tell('range_extrapolation_book', args=args)

    def add_format_setting_less_multicore(self, args):
        from .ZeroExExcelFormatter import add_result_format_setting

        add_result_format_setting(self.parent.excel_client, args, self.logger)

    def save( self, xlsx_file ):
        save_allowing_user_reply( self.wb, xlsx_file )
