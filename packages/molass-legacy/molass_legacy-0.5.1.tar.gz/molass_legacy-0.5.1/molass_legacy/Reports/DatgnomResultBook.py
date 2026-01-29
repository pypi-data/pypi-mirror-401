"""

    DatgnomResultBook.py

    Copyright (c) 2018-2023, SAXS Team, KEK-PF

"""
import os
import numpy                as np
import csv
import logging
from datetime               import datetime
# from openpyxl               import Workbook
from openpyxl.styles        import Alignment
from openpyxl.utils         import get_column_letter
from openpyxl.chart         import ScatterChart, Reference, Series
from openpyxl.chart.data_source     import StrRef
from openpyxl.chart.layout  import Layout, ManualLayout
from OpenPyXlUtil           import save_allowing_user_reply
from molass_legacy._MOLASS.Version                import get_version_string
from molass_legacy._MOLASS.SerialSettings         import get_setting
from LinearityScore         import linearity_score100, stderror_score100, FACTOR_WEIGHT
from molass_legacy.Reports.DefaultFont import set_default_font
set_default_font()

CHART_WIDTH     = 25
CHART_HEIGHT    = 15
CHART_1_POS     = 'K3'
CHART_2_POS     = 'K35'

def make_datgnom_result_book(
            excel_is_available, wb, ws,
            mjn, range_type, datgnom_out_file, book_file, parent ):

    book = DatgnomResultBook( wb, ws, mjn, range_type, datgnom_out_file, parent )

    if excel_is_available:
        book.save( book_file )
        book.add_format_setting( book_file )
    else:
        _, file = os.path.split(book_file)
        parent.logger.warning( 'excel book formatting has been skipped for %s.', file )

class DatgnomResultBook:
    def __init__( self, wb, ws, mjn, range_type, datgnom_out_file, parent ):
        self.logger = logging.getLogger( __name__ )

        m, j, num_peaks_to_exec, peak_num_ranges = mjn
        self.m  = m
        self.ud = j
        self.wb = wb
        self.ws = ws
        self.parent = parent

        paren = '' if num_peaks_to_exec == 1 else '(%d)' % ( m + 1 )
        if peak_num_ranges == 1 and range_type < 5:
            side = 'Both'
        else:
            side = 'Asc' if self.ud == 0 else 'Desc'
        ws.title = '%s-side Datgnom Result%s' % ( side, paren )

        ws.append( [ None, 'Experimental Data and Fit', None, None, None, None, 'Real Space Data' ]  )
        ws.append( [ 'S', 'J EXP', 'ERROR', 'J REG', 'I REG', None, 'R', 'P(R)', 'ERROR' ] )

        exper_fit, real_space = self.read_data( datgnom_out_file )

        self.num_rows_exper_fit     = len(exper_fit)
        self.num_rows_real_space    = len(real_space)
        num_rows = max( self.num_rows_exper_fit , self.num_rows_real_space )

        for i in range( num_rows ):
            if i < len(exper_fit):
                row = exper_fit[i]
            else:
                row = [None] * 5

            row += [None]

            if i < len(real_space):
                row += real_space[i]
            else:
                row += [None] * 3

            ws.append( row )

        self.create_scattering_curve_chart()
        self.create_pr_function_chart()

    def save( self, xlsx_file ):
        save_allowing_user_reply( self.wb, xlsx_file )

    def read_data(self, datgnom_out_file):
        from molass_legacy.ATSAS.DatGnom import datgnom_read_data
        return datgnom_read_data(datgnom_out_file)

    def create_scattering_curve_chart( self ):
        c_ = ScatterChart()
        c_.title = "Datgnom generated scattering intensity"
        c_.y_axis.title = 'log( I )'
        c_.x_axis.title = 'Q(Å⁻¹)'

        min_row = 2
        max_row = 2 + self.num_rows_exper_fit
        xvalues = Reference(self.ws, min_col=1,
                                min_row=min_row + 1,    # must not include column title
                                max_row=max_row )

        for col in [ 2, 5 ]:
            values = Reference(self.ws, min_col=col, min_row=min_row, max_row=max_row)
            series = Series(values, xvalues, title_from_data=True)
            c_.series.append(series)

        c_.y_axis.scaling.logBase = 10

        c_.width    = CHART_WIDTH
        c_.height   = CHART_HEIGHT

        adjust = 0
        c_.layout = Layout(
            ManualLayout(
            x=0.07+adjust, y=0.11,
            h=0.8,  w=0.86+adjust,
            xMode="edge",
            yMode="edge",
            )
        )

        self.ws.add_chart( c_, CHART_1_POS )

    def create_pr_function_chart( self ):
        c_ = ScatterChart()
        c_.title = "Datgnom generated P(r)"
        c_.y_axis.title = 'p(r)'
        c_.x_axis.title = 'r'

        min_row = 2
        max_row = 2 + self.num_rows_real_space
        xvalues = Reference(self.ws, min_col=7,
                                min_row=min_row + 1,    # must not include column title
                                max_row=max_row )

        for col in [ 8 ]:
            values = Reference(self.ws, min_col=col, min_row=min_row, max_row=max_row)
            series = Series(values, xvalues, title_from_data=True)
            c_.series.append(series)

        c_.width    = CHART_WIDTH
        c_.height   = CHART_HEIGHT

        adjust = 0
        c_.layout = Layout(
            ManualLayout(
            x=0.07+adjust, y=0.11,
            h=0.8,  w=0.86+adjust,
            xMode="edge",
            yMode="edge",
            )
        )

        self.ws.add_chart( c_, CHART_2_POS )

    def add_format_setting( self, book_file ):
        from .DatgnomExcelFormatter import DatgnomResultArgs
        args = DatgnomResultArgs(self.ws.title, book_file)

        if self.parent.more_multicore:
            self.add_format_setting_more_multicore(args)
        else:
            self.add_format_setting_less_multicore(args)

    def add_format_setting_more_multicore( self, args ):
        self.parent.teller.tell('range_datgnom_book', args=args)

    def add_format_setting_less_multicore( self, args ):
        from .DatgnomExcelFormatter import add_datgnom_format_setting

        add_datgnom_format_setting(self.parent.excel_client, args, self.logger)
